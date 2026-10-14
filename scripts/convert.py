import json,re,sys,hashlib
from pathlib import Path
import openpyxl
source=Path(sys.argv[1]); dest=Path(sys.argv[2]); dest.mkdir(parents=True,exist_ok=True)
w=openpyxl.load_workbook(source,data_only=True); s=w['Housing']
merged={}
for area in s.merged_cells.ranges:
    for row in range(area.min_row,area.max_row+1):
        for col in range(area.min_col,area.max_col+1): merged[row,col]=s.cell(area.min_row,area.min_col).value
records=[]
for row in range(3,s.max_row+1):
    model=merged.get((row,2),s.cell(row,2).value)
    raw=str(s.cell(row,4).value or '')
    if not raw.strip(): continue
    record={'id':f'housing-{row}','sourceRow':row,'model':str(model or '').strip(),'part':str(s.cell(row,3).value or '未标注零件名称').strip(),'raw':raw,'warnings':[]}
    for key,label in [('width','宽度'),('length','长度'),('stepHeight','台阶高度'),('stepLength','台阶长度')]:
        pattern=r'(?:^|[①②③④\n])\s*'+label+r'\s*[：:]\s*([+-]?\d+(?:\.\d+)?)'
        found=re.search(pattern,raw)
        record[key]=float(found.group(1)) if found else None
        if found is None: record['warnings'].append(label+'未提供数字')
    if record['stepHeight'] is not None and abs(record['stepHeight'])>15:
        record['warnings'].append('台阶高度异常，疑似与台阶长度填反；保留原值，请核对')
    if record['stepLength'] is not None and record['stepLength']>record['length']:
        record['warnings'].append('台阶长度大于外形长度，请核对原文的测量定义')
    records.append(record)
assert len(records)==141
assert all(r['length'] and r['width'] and r['model'] for r in records)
data={'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sheet':'Housing','units':'mm','note':s.cell(1,1).value,'records':records}
(dest/'housing-data.js').write_text('globalThis.HOUSING_DATA = '+json.dumps(data,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
print(json.dumps({'records':len(records),'flagged':sum(bool(r['warnings']) for r in records)},ensure_ascii=True))
